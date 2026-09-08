"""Builds oneday-06-decree-slips.docx — A4 sheet of 8 cut-out decree slips (editable)."""
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ACCENT, TEXT, MUTED, CUT = RGBColor(0xF5, 0x6A, 0x6A), RGBColor(0x3D, 0x44, 0x49), RGBColor(0x7F, 0x88, 0x8F), "9AA3AB"


def set_run(r, size, color, bold=False, italic=False):
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color


def dashed_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), CUT)
        borders.append(el)
    tcPr.append(borders)


def ruled_line(cell):
    # Underlined NBSPs: renders as a writing rule in Word, print, and PDF with no border-merge quirks
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("\u00A0" * 96)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.underline = True
    r.font.color.rgb = RGBColor(0xB9, 0xC0, 0xC6)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(10)

table = doc.add_table(rows=4, cols=2)
table.autofit = False

categories = "school rules  \u00b7  food  \u00b7  weekends  \u00b7  technology  \u00b7  sports  \u00b7  shows & games"
laws = "Summit law:  no naming classmates  \u00b7  nothing about anyone\u2019s body or identity"

for row in table.rows:
    row.height = Mm(68)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    for cell in row.cells:
        cell.width = Mm(95)
        dashed_borders(cell)
        p0 = cell.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        set_run(p0.add_run("DECREE PROPOSAL  \u00b7  THE DELEGATES\u2019 OWN LAWS"), 8, ACCENT, bold=True)
        p1 = cell.add_paragraph()
        p1.paragraph_format.space_after = Pt(4)
        set_run(p1.add_run("Write ONE sentence the whole summit must take a stand on:"), 9.5, TEXT)
        for _ in range(3):
            ruled_line(cell)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(10)
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(categories), 7.5, MUTED, italic=True)
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(2)
        p3.paragraph_format.space_after = Pt(0)
        set_run(p3.add_run(laws), 7.5, ACCENT)

doc.core_properties.title = "Silent Summit \u2014 Decree Slips (Grade 7 Debate Walk)"
doc.save(r"F:\Antigravity\simroom\Github Repos\bihi2027\ILT\oneday-06-decree-slips.docx")
print("saved decree slips")
