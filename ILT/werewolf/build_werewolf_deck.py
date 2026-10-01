"""Builds werewolf-ilt.pptx — Werewolf / Mafia class-game deck (16:9, dark night theme)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

W, H = 13.333, 7.5
BG     = RGBColor(0x0F, 0x17, 0x26)   # deep night navy (dominant)
PANEL  = RGBColor(0x1A, 0x25, 0x40)   # card navy
PANEL2 = RGBColor(0x14, 0x1D, 0x33)
HAIR   = RGBColor(0x2A, 0x35, 0x52)
TEXT   = RGBColor(0xF2, 0xEF, 0xE6)   # warm off-white
MUTED  = RGBColor(0x9A, 0xA6, 0xBF)
GOLD   = RGBColor(0xF2, 0xC1, 0x4E)   # moonlight (primary)
GOLDD  = RGBColor(0x6E, 0x5A, 0x24)
RED    = RGBColor(0xE4, 0x57, 0x2E)   # wolf accent (sparingly)
REDP   = RGBColor(0x35, 0x1B, 0x1E)   # dark red panel
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
FH, FB = "Arial Black", "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]

STARS = [(0.8,0.7),(2.1,1.5),(3.4,0.5),(4.9,1.2),(6.3,0.6),(7.8,1.6),(9.2,0.5),(11.0,0.9),
         (12.4,1.6),(1.5,2.4),(11.8,2.8),(0.7,1.9),(10.3,2.1),(5.7,2.3),(12.8,0.6)]

def slide(hidden=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if hidden:
        s._element.set("show", "0")
    return s

def tx(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("space_after"): para.space_after = Pt(p["space_after"])
        if p.get("line"): para.line_spacing = p["line"]
        runs = p["runs"] if "runs" in p else [p]
        for r in runs:
            run = para.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("f", FB)
            f.size = Pt(r.get("s", 17))
            f.bold = r.get("b", False)
            f.italic = r.get("i", False)
            f.color.rgb = r.get("c", TEXT)
            if r.get("spc"): f._rPr.set("spc", str(r["spc"]))
    return tb

def box(s, shape, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, rot=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    if rot: sp.rotation = rot
    sp.shadow.inherit = False
    return sp

def shape_text(sp, paras, anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.CENTER)
        if p.get("space_after"): para.space_after = Pt(p["space_after"])
        run = para.add_run(); run.text = p["t"]
        f = run.font
        f.name = p.get("f", FH); f.size = Pt(p.get("s", 14))
        f.bold = p.get("b", True); f.italic = p.get("i", False)
        f.color.rgb = p.get("c", TEXT)
        if p.get("spc"): f._rPr.set("spc", str(p["spc"]))

def kicker(s, x, y, text, color=GOLD, w=9.0, align=PP_ALIGN.LEFT):
    tx(s, x, y, w, 0.3, [{"t": text, "f": FB, "s": 12, "b": True, "c": color, "spc": 300, "align": align}])

def hair(s, x, y, w, color=HAIR):
    box(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.014, fill=color)

def pagenum(s, n):
    tx(s, W-1.1, H-0.52, 0.6, 0.3, [{"t": f"{n:02d}", "s": 12, "c": MUTED, "f": FB, "align": PP_ALIGN.RIGHT}])

def moon(s, cx, cy, d, crescent=False, on=BG):
    box(s, MSO_SHAPE.OVAL, cx-d/2, cy-d/2, d, d, fill=GOLD)
    if crescent:
        box(s, MSO_SHAPE.OVAL, cx-d/2+d*0.28, cy-d/2-d*0.16, d, d, fill=on)

def stars(s, region=(0.4, 0.4, 12.5, 2.4)):
    for i, (fx, fy) in enumerate(STARS):
        x = region[0] + fx * region[2] / 13.0
        y = region[1] + fy * region[3] / 3.0
        d = 0.035 if i % 3 else 0.055
        box(s, MSO_SHAPE.OVAL, x, y, d, d, fill=RGBColor(0x4A, 0x58, 0x78))

# ---- role icons (native shapes) -------------------------------------------
def icon_paw(s, cx, cy, k, main=RED):
    box(s, MSO_SHAPE.OVAL, cx-0.34*k, cy-0.12*k, 0.68*k, 0.56*k, fill=main)
    for dx, dy, dd in [(-0.44,-0.34,0.24),(-0.16,-0.52,0.26),(0.16,-0.52,0.26),(0.44,-0.34,0.24)]:
        box(s, MSO_SHAPE.OVAL, cx+dx*k-dd*k/2, cy+dy*k-dd*k/2, dd*k, dd*k, fill=main)

def icon_lens(s, cx, cy, k, main=GOLD):
    box(s, MSO_SHAPE.OVAL, cx-0.34*k, cy-0.44*k, 0.62*k, 0.62*k, fill=None, line=main, lw=5.5*k)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx+0.16*k, cy+0.12*k, 0.5*k, 0.16*k, fill=main, rot=45, radius=0.5)

def icon_cross(s, cx, cy, k, main=GOLD):
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx-0.11*k, cy-0.42*k, 0.22*k, 0.84*k, fill=main, radius=0.3)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx-0.42*k, cy-0.11*k, 0.84*k, 0.22*k, fill=main, radius=0.3)

def icon_house(s, cx, cy, k, main=GOLD):
    box(s, MSO_SHAPE.ISOSCELES_TRIANGLE, cx-0.5*k, cy-0.48*k, 1.0*k, 0.48*k, fill=main)
    box(s, MSO_SHAPE.RECTANGLE, cx-0.36*k, cy-0.02*k, 0.72*k, 0.48*k, fill=main)

ICONS = {"wolf": icon_paw, "detective": icon_lens, "doctor": icon_cross, "villager": icon_house}

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

# ============================ 1 · TITLE =====================================
s = slide()
stars(s)
moon(s, 10.9, 2.5, 2.6, crescent=True)
kicker(s, 0.9, 2.05, "INTEGRATED LEARNING TIME  ·  CLASS GAME")
tx(s, 0.87, 2.35, 9.5, 1.7, [{"runs": [
    {"t": "WERE", "f": FH, "s": 88, "b": True, "c": TEXT},
    {"t": "WOLF", "f": FH, "s": 88, "b": True, "c": GOLD}]}])
tx(s, 0.9, 4.15, 8.6, 0.9, [{"t": "A social deduction game of secrets, suspicion, and smart voting.",
    "s": 20, "c": MUTED, "line": 1.15}])
hair(s, 0.9, 5.45, 7.2)
tx(s, 0.9, 5.7, 11.5, 0.5, [{"runs": [
    {"t": "28 players", "s": 15, "b": True, "c": TEXT}, {"t": "   ·   ", "s": 15, "c": MUTED},
    {"t": "1 moderator", "s": 15, "b": True, "c": TEXT}, {"t": "   ·   ", "s": 15, "c": MUTED},
    {"t": "2 hidden teams", "s": 15, "b": True, "c": TEXT}, {"t": "   ·   ", "s": 15, "c": MUTED},
    {"t": "45–60 minutes", "s": 15, "b": True, "c": TEXT}]}])
notes(s, "Open here. Say: 'We're playing Werewolf today — also called Mafia. It's a game of secrets and arguments. Some of you will be given a secret role. Nobody knows who is who — except the wolves. I run the game; I am NOT a player.'")

# ============================ 2 · THE STORY =================================
s = slide()
kicker(s, 0.6, 0.55, "THE SET-UP")
tx(s, 0.57, 0.9, 6.6, 1.0, [{"t": "Wolves are hiding in the village.", "f": FH, "s": 34, "b": True}])
tx(s, 0.6, 2.15, 6.4, 3.4, [
    {"t": "Everyone in this room is a villager. But three of you are secretly werewolves — and only they know who their pack is.", "s": 17, "c": TEXT, "line": 1.25, "space_after": 12},
    {"t": "Every night, the wolves silently choose one victim. Every day, the village wakes, discovers the loss — and has five minutes to work out who is lying.", "s": 17, "c": TEXT, "line": 1.25, "space_after": 12},
    {"t": "The village argues. Accuses. Votes. And the wolves vote right along with everyone else.", "s": 17, "c": MUTED, "i": True, "line": 1.25}])
# right: night/day cycle graphic
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 7.6, 0.9, 5.1, 5.7, fill=PANEL, radius=0.06)
moon(s, 10.15, 2.15, 1.5)
tx(s, 8.1, 3.05, 4.1, 0.35, [{"t": "NIGHT — the wolves hunt", "f": FH, "s": 15, "b": True, "c": GOLD, "align": PP_ALIGN.CENTER}])
tx(s, 8.1, 3.42, 4.1, 0.55, [{"t": "Eyes closed. Silence. Victims fall.", "s": 13.5, "c": MUTED, "align": PP_ALIGN.CENTER}])
box(s, MSO_SHAPE.DOWN_ARROW, 9.95, 4.05, 0.4, 0.5, fill=HAIR)
box(s, MSO_SHAPE.OVAL, 9.55, 4.65, 1.2, 1.2, fill=GOLD)
tx(s, 8.1, 5.95, 4.1, 0.35, [{"t": "DAY — the village votes back", "f": FH, "s": 15, "b": True, "c": TEXT, "align": PP_ALIGN.CENTER}])
pagenum(s, 2)
notes(s, "Keep this to 60 seconds of storytelling. Drama sells the game: lower your voice on 'every night, the wolves take a victim.'")

# ============================ 3 · HOW YOU WIN ===============================
s = slide()
kicker(s, 0.6, 0.55, "TWO TEAMS · TWO WAYS TO WIN")
tx(s, 0.57, 0.9, 9.0, 0.9, [{"t": "How you win", "f": FH, "s": 34, "b": True}])
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 2.0, 5.9, 3.4, fill=PANEL, radius=0.05)
icon_house(s, 1.55, 2.95, 0.9)
tx(s, 2.5, 2.55, 3.8, 0.9, [{"t": "VILLAGE WINS", "f": FH, "s": 21, "b": True, "c": GOLD},
                             {"t": "when the last wolf is voted out.", "s": 15, "c": TEXT, "b": False}])
tx(s, 1.0, 4.0, 5.1, 1.2, [{"t": "Every single werewolf has been eliminated by day votes.", "s": 15, "c": MUTED, "line": 1.25}])
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 6.85, 2.0, 5.9, 3.4, fill=REDP, radius=0.05)
icon_paw(s, 7.85, 2.95, 0.75)
tx(s, 8.8, 2.55, 3.8, 0.9, [{"t": "WOLVES WIN", "f": FH, "s": 21, "b": True, "c": RED},
                             {"t": "when wolves match the villagers.", "s": 15, "c": TEXT, "b": False}])
tx(s, 7.25, 4.0, 5.1, 1.2, [{"t": "Wolves equal or outnumber the living villagers — the night feed is unstoppable.", "s": 15, "c": MUTED, "line": 1.25}])
box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.85, 12.15, 0.85, fill=PANEL2, radius=0.14)
tx(s, 1.0, 6.08, 11.4, 0.5, [{"runs": [
    {"t": "Eliminated players are not out of the fun — ", "s": 15, "c": TEXT},
    {"t": "they become silent spectators", "s": 15, "b": True, "c": GOLD},
    {"t": " and watch every secret unfold at the reveal.", "s": 15, "c": TEXT}]}])
pagenum(s, 3)
notes(s, "Stress the spectator rule early — it defuses 'I'm out, this is boring.' Eliminated players stay in the room, watch everything, and get their questions answered at the final reveal.")

# ============================ 4 · WHO'S WHO =================================
s = slide()
kicker(s, 0.6, 0.55, "SECRET ROLES · STANDARD GAME · 28 PLAYERS")
tx(s, 0.57, 0.9, 9.0, 0.9, [{"t": "Who's who", "f": FH, "s": 34, "b": True}])
cards = [
    ("wolf", "WEREWOLF", "×3", "Wakes with the pack and picks tonight's victim.", RED, REDP),
    ("detective", "DETECTIVE", "×1", "Checks one player each night: wolf or innocent?", GOLD, PANEL),
    ("doctor", "DOCTOR", "×1", "Shields one player each night from the wolves.", GOLD, PANEL),
    ("villager", "VILLAGER", "×23", "No powers. Questions, listening — and a vote.", GOLD, PANEL),
]
cw, gap = 2.94, 0.13
for i, (ic, name, count, desc, col, pfill) in enumerate(cards):
    x = 0.6 + i * (cw + gap)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, cw, 4.35, fill=pfill, radius=0.05)
    ICONS[ic](s, x + cw/2, 3.0, 0.85, main=col if ic != "wolf" else RED)
    tx(s, x+0.2, 3.75, cw-0.4, 0.4, [{"t": name, "f": FH, "s": 17, "b": True, "c": TEXT, "align": PP_ALIGN.CENTER}])
    tx(s, x+0.2, 4.18, cw-0.4, 0.5, [{"t": count + " in play", "f": FH, "s": 15, "b": True,
                                       "c": col, "align": PP_ALIGN.CENTER}])
    tx(s, x+0.28, 4.75, cw-0.56, 1.4, [{"t": desc, "s": 13.5, "c": MUTED, "line": 1.2, "align": PP_ALIGN.CENTER}])
tx(s, 0.6, 6.65, 12.1, 0.4, [{"t": "Your role card is secret. Keep it hidden — even from your best friend. Especially from your best friend.",
    "s": 14, "i": True, "c": MUTED, "align": PP_ALIGN.CENTER}])
pagenum(s, 4)
notes(s, "If you run the 4-wolf variant with a strong group: swap one villager card for the spare wolf card — 4 wolves, 22 villagers.")

# ======================= 5–8 · ROLE SLIDES ==================================
def role_slide(icon, name, team, teamcol, panel_fill, rows, tip, count):
    s = slide()
    kicker(s, 0.6, 0.55, f"YOUR SECRET ROLE  ·  {count}", color=teamcol)
    tx(s, 0.57, 0.9, 7.3, 1.0, [{"t": name, "f": FH, "s": 40, "b": True, "c": TEXT}])
    y = 2.25
    for label, body in rows:
        tx(s, 0.6, y, 1.35, 0.35, [{"t": label, "f": FB, "s": 13, "b": True, "c": teamcol, "spc": 200}])
        tx(s, 2.05, y-0.04, 5.85, 1.15, [{"t": body, "s": 16.5, "c": TEXT, "line": 1.2}])
        y += 1.42
        if label != rows[-1][0]: hair(s, 0.6, y-0.28, 7.3)
    p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.25, 0.95, 4.5, 5.65, fill=panel_fill, radius=0.06)
    ICONS[icon](s, 10.5, 2.55, 1.35)
    tx(s, 8.65, 4.05, 3.7, 0.3, [{"t": "TIP", "f": FB, "s": 12, "b": True, "c": teamcol, "spc": 300}])
    tx(s, 8.65, 4.42, 3.75, 1.9, [{"t": tip, "s": 15, "i": True, "c": TEXT, "line": 1.3}])
    pagenum(s, 0)
    return s

s = role_slide("wolf", "Werewolf", "TEAM: WOLVES", RED, REDP, [
    ("NIGHT", "Wake silently with the other wolves. Agree — by pointing, never speaking — on one villager to eliminate."),
    ("DAY", "Blend in. Accuse someone else. Lie if you must, but never get caught in a lie.")],
    "You know your pack from night one. Protect them, deflect suspicion, and let loud villagers do your work for you.",
    "3 OF 28 PLAYERS · TEAM: WOLVES")
pagenum(s, 5)
notes(s, "Read only the NIGHT and DAY lines aloud if revealing roles this way. Say to wolves privately via card text: 'You win by surviving votes.'")

s = role_slide("detective", "The Detective", "TEAM: VILLAGE", GOLD, PANEL, [
    ("NIGHT", "Wake alone. Point at any player. The moderator silently answers: thumbs up = WEREWOLF, thumbs down = innocent."),
    ("DAY", "You hold real information. Guard your secret — reveal it too early and you'll be dead before dawn.")],
    "Collect two or three certain answers before you stake your reputation on one. A detective who dies is a village without a map.",
    "1 OF 28 PLAYERS · TEAM: VILLAGE")
pagenum(s, 6)
notes(s, "Practise the silent signal BEFORE the first night: thumbs up = wolf, thumbs down = innocent. Keep your hand low, near your chest, so a peeking neighbour can't see it.")

s = role_slide("doctor", "The Doctor", "TEAM: VILLAGE", GOLD, PANEL, [
    ("NIGHT", "Wake alone. Point at one player to protect. If the wolves attacked them tonight, they survive and nobody dies."),
    ("DAY", "Say nothing. A doctor who brags becomes the wolves' second target — right after the detective.")],
    "You MAY protect yourself. You may NOT guard the same person two nights running.",
    "1 OF 28 PLAYERS · TEAM: VILLAGE")
pagenum(s, 7)
notes(s, "If the doctor protects the wolves' target, announce at dawn: 'The wolves attacked last night — but the doctor was there. Nobody died.' That moment always gets a big reaction; milk it.")

s = role_slide("villager", "The Villager", "TEAM: VILLAGE", GOLD, PANEL, [
    ("NIGHT", "Sleep. You have no powers — which is exactly why the village needs your clear head by day."),
    ("DAY", "Your weapons are questions, listening, and your vote. Track who changes their story.")],
    "The loudest voice in the room is not automatically the wolf. Wolves are loud on purpose.",
    "23 OF 28 PLAYERS · TEAM: VILLAGE")
pagenum(s, 8)
notes(s, "Reassure villagers: having no power isn't boring — the villagers hold the only real weapon in the game, the vote.")

# ============================ 9 · THE LOOP ==================================
s = slide()
kicker(s, 0.6, 0.55, "GAME FLOW")
tx(s, 0.57, 0.9, 10.0, 0.9, [{"t": "One round = one night + one day", "f": FH, "s": 32, "b": True}])
steps = [
    ("1", "NIGHT · 90 sec", "Eyes closed. Wolves hunt, detective investigates, doctor protects."),
    ("2", "DAWN", "Moderator announces who died — or that everyone survived."),
    ("3", "DEBATE · 5 min", "Ask questions. Make accusations. Defend yourself."),
    ("4", "THE VOTE", "Anonymous ballot slip. Majority is eliminated."),
    ("5", "CHECK", "All wolves gone? Village wins. Wolves = villagers? Wolves win."),
]
bw, bgap = 2.25, 0.28
x0 = (W - (5*bw + 4*bgap)) / 2
for i, (num, head, body) in enumerate(steps):
    x = x0 + i*(bw+bgap)
    fill = REDP if i == 0 else PANEL
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.35, bw, 3.3, fill=fill, radius=0.07)
    c = box(s, MSO_SHAPE.OVAL, x+0.18, 2.55, 0.5, 0.5, fill=GOLD if i else RED)
    shape_text(c, [{"t": num, "s": 18, "c": BG if i else WHITE}])
    tx(s, x+0.18, 3.25, bw-0.36, 0.6, [{"t": head, "f": FH, "s": 14.5, "b": True, "c": GOLD if i else RED, "line": 1.05}])
    tx(s, x+0.18, 3.95, bw-0.36, 1.6, [{"t": body, "s": 13, "c": TEXT, "line": 1.22}])
    if i < 4:
        box(s, MSO_SHAPE.RIGHT_ARROW, x+bw+0.03, 3.75, 0.22, 0.3, fill=MUTED)
arr = box(s, MSO_SHAPE.LEFT_ARROW, x0+0.4, 6.05, 5*bw+4*bgap-0.8, 0.55, fill=PANEL2)
shape_text(arr, [{"t": "REPEAT — each round takes 8–10 minutes · play 3–5 rounds", "s": 13, "c": MUTED, "f": FB}])
pagenum(s, 9)
notes(s, "This is the whole game on one slide. Come back to it between rounds if anyone looks lost.")

# ============================ 10 · NIGHT RULES ==============================
s = slide()
kicker(s, 0.6, 0.55, "PHASE ONE")
tx(s, 0.57, 0.9, 8.0, 0.9, [{"t": "Night — eyes closed", "f": FH, "s": 34, "b": True}])
rules = [
    ("Heads down, eyes closed, absolute silence.", "The game only works if nobody peeks and nobody talks."),
    ("The moderator calls each role in order.", "Wolves first, then the detective, then the doctor. Everyone else stays asleep."),
    ("Act by pointing — never by sound.", "No whispers, no giggling, no chair-creaking tells."),
    ("Peeking once = a warning. Twice = you sit out.", "Harsh? Yes. Fair to 27 other players? Also yes."),
]
y = 2.15
for i, (head, body) in enumerate(rules):
    c = box(s, MSO_SHAPE.OVAL, 0.6, y+0.03, 0.46, 0.46, fill=GOLD)
    shape_text(c, [{"t": str(i+1), "s": 16, "c": BG}])
    tx(s, 1.35, y, 6.7, 0.45, [{"t": head, "f": FB, "s": 16.5, "b": True}])
    tx(s, 1.35, y+0.48, 6.7, 0.55, [{"t": body, "s": 14, "c": MUTED, "line": 1.15}])
    y += 1.18
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.6, 2.0, 4.15, 4.6, fill=PANEL, radius=0.06)
moon(s, 10.67, 3.5, 1.7, crescent=True, on=PANEL)
tx(s, 8.9, 4.75, 3.55, 0.5, [{"t": "SILENCE", "f": FH, "s": 30, "b": True, "c": GOLD, "align": PP_ALIGN.CENTER, "spc": 400}])
tx(s, 8.9, 5.35, 3.55, 0.9, [{"t": "If you talk at night, you steal the game from everyone else.",
    "s": 13.5, "i": True, "c": MUTED, "align": PP_ALIGN.CENTER, "line": 1.25}])
pagenum(s, 10)
notes(s, "Script: 'Night falls. Everyone close your eyes — heads down on the desk. NO PEEKING.' Then the wake order: wolves, detective, doctor. Use the exact lines on your pocket card.")

# ============================ 11 · DAY RULES ================================
s = slide()
kicker(s, 0.6, 0.55, "PHASE TWO")
tx(s, 0.57, 0.9, 9.0, 0.9, [{"t": "Day — the debate", "f": FH, "s": 34, "b": True}])
steps = [
    ("The moderator announces the night's victim.", "That player becomes a silent spectator immediately."),
    ("60 seconds of silence first.", "Think. Look around. Let the shock settle before anyone speaks."),
    ("Open debate — five minutes, hard cap.", "Speak only while holding the talking object. The timer rules."),
    ("Every accusation needs a BECAUSE.", "\u201cI suspect Sam because...\u201d No because = no floor time."),
]
y = 2.15
for i, (head, body) in enumerate(steps):
    c = box(s, MSO_SHAPE.OVAL, 0.6, y+0.03, 0.46, 0.46, fill=GOLD)
    shape_text(c, [{"t": str(i+1), "s": 16, "c": BG}])
    tx(s, 1.35, y, 7.0, 0.45, [{"t": head, "f": FB, "s": 16.5, "b": True}])
    tx(s, 1.35, y+0.48, 7.0, 0.55, [{"t": body, "s": 14, "c": MUTED, "line": 1.15}])
    y += 1.18
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.85, 2.0, 3.9, 4.6, fill=PANEL, radius=0.06)
box(s, MSO_SHAPE.OVAL, 10.15, 2.5, 1.3, 1.3, fill=GOLD)
import math
for adeg in range(0, 360, 45):
    rr = 0.95
    cx = 10.8 + rr*math.cos(math.radians(adeg)) - 0.05
    cy = 3.15 + rr*math.sin(math.radians(adeg)) - 0.05
    box(s, MSO_SHAPE.OVAL, cx, cy, 0.1, 0.1, fill=GOLD)
tx(s, 9.15, 4.35, 3.3, 0.5, [{"t": "5:00", "f": FH, "s": 34, "b": True, "c": TEXT, "align": PP_ALIGN.CENTER}])
tx(s, 9.15, 5.05, 3.3, 1.3, [{"t": "Debate cap per day. The moderator's phone timer is the law.",
    "s": 13.5, "i": True, "c": MUTED, "align": PP_ALIGN.CENTER, "line": 1.25}])
pagenum(s, 11)
notes(s, "Use a visible timer. When it hits zero, say: 'Time. Hands on heads — vote.' A hard cap keeps the pace and forces decisions.")

# ============================ 12 · VOTING ===================================
s = slide()
kicker(s, 0.6, 0.55, "THE DAILY VOTE")
tx(s, 0.57, 0.9, 9.0, 0.9, [{"t": "Votes are silent. And secret.", "f": FH, "s": 34, "b": True}])
vsteps = [
    ("Write one name on your ballot slip.", "Anyone except yourself. Fold it. Drop it in the box."),
    ("The moderator counts aloud.", "Every vote is read out loud so the count is public."),
    ("Majority = eliminated. Tie = nobody dies.", "A tied village wakes everyone again the next night."),
    ("Last words, then silence.", "The eliminated player gets ONE sentence, then spectates quietly."),
]
y = 2.15
for i, (head, body) in enumerate(vsteps):
    c = box(s, MSO_SHAPE.OVAL, 0.6, y+0.03, 0.46, 0.46, fill=GOLD)
    shape_text(c, [{"t": str(i+1), "s": 16, "c": BG}])
    tx(s, 1.35, y, 7.0, 0.45, [{"t": head, "f": FB, "s": 16.5, "b": True}])
    tx(s, 1.35, y+0.48, 7.0, 0.55, [{"t": body, "s": 14, "c": MUTED, "line": 1.15}])
    y += 1.18
# ballot box graphic
p = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 8.85, 2.0, 3.9, 4.6, fill=PANEL, radius=0.06)
for i, (sx, sy, rot) in enumerate([(10.0, 2.35, -14), (10.55, 2.5, 9), (9.7, 2.75, 4)]):
    box(s, MSO_SHAPE.RECTANGLE, sx, sy, 0.85, 0.55, fill=TEXT, rot=rot)
bxs = box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.45, 3.45, 1.7, 1.35, fill=GOLD, radius=0.08)
box(s, MSO_SHAPE.RECTANGLE, 10.05, 3.35, 0.5, 0.12, fill=PANEL)
tx(s, 9.05, 5.05, 3.5, 0.5, [{"t": "Anonymous = fearless", "f": FH, "s": 16, "b": True, "c": GOLD, "align": PP_ALIGN.CENTER}])
tx(s, 9.05, 5.6, 3.5, 0.9, [{"t": "Quiet players vote as hard as loud ones.",
    "s": 13.5, "i": True, "c": MUTED, "align": PP_ALIGN.CENTER, "line": 1.25}])
pagenum(s, 12)
notes(s, "Anonymous slips beat hands-up voting for shy classes. Have a box or mug ready, plus printed ballot slips (in your print pack).")

# ============================ 13 · STRATEGY =================================
s = slide()
kicker(s, 0.6, 0.55, "LEVEL UP")
tx(s, 0.57, 0.9, 11.0, 0.9, [{"t": "How to actually be good at this", "f": FH, "s": 32, "b": True}])
cols = [
    ("IF YOU'RE A VILLAGER", GOLD, PANEL, [
        "Ask quiet players directly — wolves love a debate where only loud people talk.",
        "Compare stories: who said something yesterday that doesn't match today?",
        "Vote with the group's evidence, not with its volume."]),
    ("IF YOU'RE A WOLF", RED, REDP, [
        "Act like a villager: agree a little, accuse slowly, never defend your pack too hard.",
        "Get one innocent voted out early and the village will trust your judgement.",
        "If a packmate is doomed, join the vote against them — sacrifice keeps you alive."]),
    ("IF YOU'RE SPECIAL", GOLD, PANEL, [
        "Detective: never reveal on day one. Survive first, inform later.",
        "Doctor: protect silently, predict where the wolves will strike.",
        "One good night of information beats three days of guessing."]),
]
cw = 3.93
for i, (head, col, pfill, items) in enumerate(cols):
    x = 0.6 + i*(cw+0.18)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.0, cw, 4.55, fill=pfill, radius=0.05)
    tx(s, x+0.3, 2.3, cw-0.6, 0.4, [{"t": head, "f": FB, "s": 15, "b": True, "c": col, "spc": 100}])
    hair(s, x+0.3, 2.85, cw-0.6, color=col if col == RED else HAIR)
    paras = []
    for it in items:
        paras.append({"runs": [{"t": "—  ", "s": 14.5, "c": col, "b": True},
                               {"t": it, "s": 14.5, "c": TEXT}], "line": 1.25, "space_after": 12})
    tx(s, x+0.3, 3.1, cw-0.6, 3.2, paras)
pagenum(s, 13)
notes(s, "Don't lecture this slide — flash it for 30 seconds. Half the fun is kids discovering these strategies themselves.")

# ============================ 14 · VILLAGE CODE =============================
s = slide()
kicker(s, 0.6, 0.55, "FAIR PLAY")
tx(s, 0.57, 0.9, 10.0, 0.9, [{"t": "The village code", "f": FH, "s": 34, "b": True}])
codes = [
    ("No peeking at night. Ever.", "One warning, then you sit out. This protects everyone's fun."),
    ("Accuse the argument, not the person.", "No real-world jabs, no gangs, no grudges from outside the game."),
    ("Losing is part of the game.", "A brilliant wolf win deserves applause as loud as a village win."),
    ("Spectators are shadows.", "Once you're out: watch everything, say nothing — no faces, no hints."),
]
y = 2.2
for i, (head, body) in enumerate(codes):
    tx(s, 0.6, y-0.12, 1.2, 0.9, [{"t": f"{i+1:02d}", "f": FH, "s": 30, "b": True, "c": GOLDD}])
    tx(s, 1.85, y, 10.5, 0.45, [{"t": head, "f": FB, "s": 17, "b": True}])
    tx(s, 1.85, y+0.47, 10.5, 0.5, [{"t": body, "s": 14, "c": MUTED, "line": 1.15}])
    y += 1.13
    if i < 3: hair(s, 0.6, y-0.22, 12.1)
pagenum(s, 14)
notes(s, "Read rule 2 with weight — this is the rule that keeps the game safe. 'It's a game about lying TO EACH OTHER in character; the moment it's about a real person, it stops.'")

# ============================ 15 · DEBRIEF ==================================
s = slide()
kicker(s, 0.6, 0.55, "AFTER THE FINAL REVEAL · 10 MINUTES")
tx(s, 0.57, 0.9, 10.5, 0.9, [{"t": "Talk about it", "f": FH, "s": 34, "b": True}])
qs = [
    "What actually helped you decide your vote — and what just felt like noise?",
    "When the loudest voice won an argument, were they right?",
    "Did you change your mind mid-day? What changed it?",
    "Wolves: what was hardest about lying to your friends?",
    "Detective: when did you sit on information, and why?",
    "What should the village do differently in the next game?",
]
y = 2.1
for i, q in enumerate(qs):
    c = box(s, MSO_SHAPE.OVAL, 0.6, y+0.02, 0.44, 0.44, fill=PANEL, line=GOLD, lw=1.2)
    shape_text(c, [{"t": "Q" + str(i+1), "s": 12, "c": GOLD, "f": FB}])
    tx(s, 1.3, y, 11.3, 0.55, [{"t": q, "s": 16.5, "c": TEXT, "line": 1.1}])
    y += 0.79
tx(s, 1.3, 6.95, 11.0, 0.4, [{"t": "You're practising evidence, source-checking and group psychology — that's the learning, disguised as fun.",
    "s": 13, "i": True, "c": MUTED}])
pagenum(s, 15)
notes(s, "Pick 3 questions if time is short. This debrief is your ILT evidence: evidence-based reasoning, persuasion, and group dynamics — students name it themselves.")

# ============================ 16 · CLOSING ==================================
s = slide(hidden=False)
stars(s)
moon(s, 6.67, 3.0, 2.3, crescent=True)
tx(s, 0.6, 4.35, 12.1, 1.2, [{"t": "Close your eyes.", "f": FH, "s": 54, "b": True, "c": TEXT, "align": PP_ALIGN.CENTER}])
tx(s, 0.6, 5.6, 12.1, 0.6, [{"t": "Night one begins. Good luck, village.", "s": 20, "i": True, "c": MUTED, "align": PP_ALIGN.CENTER}])
notes(s, "Deal role cards face down BEFORE this slide. Then: 'Village — you are 28 strong. Wolves — I know who you are. Nobody else does. Close your eyes...'")

# ==================== 17 · TEACHER SETUP (HIDDEN) ===========================
s = slide(hidden=True)
kicker(s, 0.6, 0.55, "TEACHER SETUP — HIDDEN IN SLIDESHOW", color=RED)
tx(s, 0.57, 0.9, 10.0, 0.9, [{"t": "Before class — 10 minutes", "f": FH, "s": 32, "b": True}])
tx(s, 0.6, 2.0, 6.2, 4.6, [
    {"runs": [{"t": "DEAL THE ROLES (standard 28)", "f": FB, "s": 14, "b": True, "c": GOLD}], "space_after": 6},
    {"t": "3 Werewolf · 1 Detective · 1 Doctor · 23 Villager — shuffle and deal face down, one per student. Spares stay in your folder.", "s": 14.5, "line": 1.25, "space_after": 14},
    {"runs": [{"t": "SPICY VARIANT", "f": FB, "s": 14, "b": True, "c": GOLD}], "space_after": 6},
    {"t": "Strong group? Swap one villager for the 4th wolf card → 4 wolves, 22 villagers. Accusations fly faster.", "s": 14.5, "line": 1.25, "space_after": 14},
    {"runs": [{"t": "MATERIALS", "f": FB, "s": 14, "b": True, "c": GOLD}], "space_after": 6},
    {"t": "Role cards (cut) · ballot slips · a box or mug · a talking object (marker) · phone timer · your pocket night-order card.", "s": 14.5, "line": 1.25},
])
tx(s, 7.2, 2.0, 5.5, 4.6, [
    {"runs": [{"t": "TIMING PLAN", "f": FB, "s": 14, "b": True, "c": GOLD}], "space_after": 6},
    {"t": "30 min → 3 rounds, tight", "s": 14.5, "line": 1.3, "space_after": 4},
    {"t": "45 min → 4 rounds + short debrief", "s": 14.5, "line": 1.3, "space_after": 4},
    {"t": "60 min → 2 games + full debrief", "s": 14.5, "line": 1.3, "space_after": 14},
    {"runs": [{"t": "GOLDEN MODERATOR RULES", "f": FB, "s": 14, "b": True, "c": GOLD}], "space_after": 6},
    {"t": "You never play. You never hint. Follow the pocket card script word-for-word the first game — it keeps nights tight and fair.", "s": 14.5, "line": 1.25},
])
notes(s, "This slide is hidden — it will not show during the slideshow. Print the facilitator guide (DOCX) for the full script.")

prs.save("D:/temp/ILT/werewolf-ilt.pptx")
print("saved D:/temp/ILT/werewolf-ilt.pptx with", len(prs.slides._sldIdLst), "slides")
