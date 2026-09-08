"""Builds oneday-06-summit-deck.pptx — Grade 7 Debate Walk 'Summit of Two Kingdoms' deck."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

DARK, TEXT, MUTED = RGBColor(0x2E, 0x38, 0x3E), RGBColor(0x3D, 0x44, 0x49), RGBColor(0x7F, 0x88, 0x8F)
ACCENT, TINT, HAIR = RGBColor(0xF5, 0x6A, 0x6A), RGBColor(0xFD, 0xEE, 0xEE), RGBColor(0xE5, 0xE7, 0xE9)
PANEL, GHOST, ONDARK, WHITE = RGBColor(0xF4, 0xF5, 0xF6), RGBColor(0xFA, 0xD8, 0xD8), RGBColor(0xAE, 0xB6, 0xBD), RGBColor(0xFF, 0xFF, 0xFF)
PINK, PINK_IN, PINK_DEEP = RGBColor(0xFF, 0x9E, 0xBB), RGBColor(0xE8, 0x7A, 0x9E), RGBColor(0xD9, 0x5F, 0x8A)
FONT, W = "Arial", 13.333

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(dark=False, hidden=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = DARK if dark else WHITE
    if hidden:
        s._element.set("show", "0")
    return s


def tx(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", PP_ALIGN.LEFT)
        if "sb" in para:
            p.space_before = Pt(para["sb"])
        for t, o in para["runs"]:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name, f.size, f.bold, f.italic = FONT, Pt(o.get("size", 18)), o.get("bold", False), o.get("italic", False)
            f.color.rgb = o.get("color", TEXT)
    return tb


def box(s, x, y, w, h, fill=None, line=None, lw=1.0, kind=MSO_SHAPE.RECTANGLE, radius=None, alpha=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid() if fill else sh.fill.background()
    if fill:
        sh.fill.fore_color.rgb = fill
        if alpha is not None:
            sf = sh._element.spPr.find(qn("a:solidFill"))
            c = sf.find(qn("a:srgbClr"))
            c.append(c.makeelement(qn("a:alpha"), {"val": str(int(alpha * 1000))}))
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb, sh.line.width = line, Pt(lw)
    if radius is not None:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False
    return sh


def arrow(s, x1, y1, x2, y2, color=ACCENT, lw=2.0):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb, c.line.width = color, Pt(lw)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    c.shadow.inherit = False
    return c


def note(s, text):
    s.notes_slide.notes_text_frame.text = text


def rules_rows(s, rows, y0=1.5, rh=1.02, textw=8.6):
    for i, (head, gloss) in enumerate(rows):
        y = y0 + i * rh
        tx(s, 0.55, y - 0.03, 0.75, 0.6, [{"runs": [(str(i + 1), {"size": 26, "bold": True, "color": ACCENT})]}])
        tx(s, 1.45, y, textw, rh - 0.08, [
            {"runs": [(head, {"size": 19, "bold": True})]},
            {"runs": [(gloss, {"size": 13.5, "color": MUTED})], "sb": 3},
        ])
        if i < len(rows) - 1:
            box(s, 0.55, y + rh - 0.09, 10.4, 0.012, fill=HAIR)


def bunny(s, cx, bottom, h, glow=True):
    """Flat pink glowing plastic bunny. cx=center-x (in), bottom=body bottom (in), h=total height (in)."""
    cy = bottom - h / 2
    if glow:
        for f, a in [(1.55, 6), (1.25, 11), (1.0, 18)]:
            box(s, cx - h * f / 2, cy - h * f / 2, h * f, h * f, fill=PINK, kind=MSO_SHAPE.OVAL, alpha=a)
    for sgn in (-1, 1):
        ear = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + sgn * 0.17 * h - 0.11 * h), Inches(bottom - h * 1.02), Inches(0.22 * h), Inches(0.52 * h))
        ear.rotation = sgn * 8
        ear.fill.solid(); ear.fill.fore_color.rgb = PINK
        ear.line.fill.background(); ear.shadow.inherit = False
        inn = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + sgn * 0.17 * h - 0.05 * h), Inches(bottom - h * 0.94), Inches(0.10 * h), Inches(0.32 * h))
        inn.rotation = sgn * 8
        inn.fill.solid(); inn.fill.fore_color.rgb = PINK_IN
        inn.line.fill.background(); inn.shadow.inherit = False
    box(s, cx - 0.475 * h, bottom - 0.55 * h, 0.95 * h, 0.55 * h, fill=PINK, kind=MSO_SHAPE.OVAL)
    box(s, cx - 0.31 * h, bottom - h * 0.86, 0.62 * h, 0.58 * h, fill=PINK, kind=MSO_SHAPE.OVAL)
    for sgn in (-1, 1):
        box(s, cx + sgn * 0.11 * h - 0.03 * h, bottom - h * 0.62, 0.06 * h, 0.06 * h, fill=DARK, kind=MSO_SHAPE.OVAL)
    box(s, cx - 0.025 * h, bottom - h * 0.51, 0.05 * h, 0.05 * h, fill=PINK_DEEP, kind=MSO_SHAPE.OVAL)


def decree(numeral, tag, statement, footer, hidden=False, letter=None):
    s = slide(hidden=hidden)
    tx(s, 0.62, 0.6, 8.5, 0.4, [{"runs": [(tag, {"size": 15, "bold": True, "color": ACCENT})]}])
    tx(s, 0.3, 0.7, 3.1, 4.7, [{"runs": [(letter or numeral, {"size": 175 if len(letter or numeral) < 3 else 120, "bold": True, "color": ACCENT})]}],
       anchor=MSO_ANCHOR.MIDDLE)
    tx(s, 3.75, 1.55, 9.0, 2.8, [{"runs": [("\u201c" + statement + "\u201d", {"size": 40, "bold": True})]}])
    tx(s, 3.78, 6.55, 8.9, 0.4, [{"runs": [(footer, {"size": 12.5, "color": MUTED})]}])
    return s


FOOT = "Take your stand  \u00b7  fold the line  \u00b7  defect with honour"

# ── 1 · Title ────────────────────────────────────────────────────────────────
s = slide(dark=True)
tx(s, 8.75, -0.75, 4.3, 4.6, [{"runs": [("\u201d", {"size": 300, "bold": True, "color": RGBColor(0x3A, 0x45, 0x52)})]}])
tx(s, 0.9, 0.9, 9.0, 0.4, [{"runs": [("INTEGRATED LEARNING TIME  \u00b7  GRADE 7", {"size": 13, "bold": True, "color": ONDARK})]}])
tx(s, 0.86, 2.0, 11.9, 1.9, [{"runs": [("THE SUMMIT OF", {"size": 44, "bold": True, "color": WHITE})]},
                             {"runs": [("TWO KINGDOMS", {"size": 60, "bold": True, "color": WHITE}),
                                       (".", {"size": 60, "bold": True, "color": ACCENT})], "sb": 2}])
tx(s, 0.9, 4.55, 10.5, 0.6, [{"runs": [("One hour. 28 delegates. A thousand-year quarrel.", {"size": 24, "bold": True, "color": ACCENT})]}])
tx(s, 0.9, 5.35, 10.5, 0.5, [{"runs": [("YeahTotally  vs.  NoWay  \u2014  and the Mid in between.", {"size": 16, "color": ONDARK})]}])
note(s, "Put this on the screen as they walk in. Props off to the side: two wall signs, a pink glowing plastic bunny, "
        "a sealed envelope. Do not explain anything yet.")

# ── 2 · Cold open 1 ──────────────────────────────────────────────────────────
s = slide(dark=True)
tx(s, 1.2, 2.5, 10.9, 2.2, [
    {"align": PP_ALIGN.CENTER, "runs": [("A thousand years ago,", {"size": 40, "bold": True, "color": WHITE})]},
    {"align": PP_ALIGN.CENTER, "runs": [("there was one country.", {"size": 40, "bold": True, "color": ONDARK})], "sb": 6},
])
note(s, "COLD OPEN part 1 of 4 \u2014 read the briefing deadpan across slides 2\u20135, 45 seconds total: "
        "\u201cDelegates \u2014 a thousand years ago, YeahTotally and NoWay were one country.\u201d")

# ── 3 · Cold open 2 — THE LOUD GOD ──────────────────────────────────────────
s = slide(dark=True)
tx(s, 0.9, 1.7, 7.3, 0.4, [{"runs": [("THE ENEMY", {"size": 13, "bold": True, "color": ONDARK})]}])
tx(s, 0.86, 2.2, 7.6, 1.3, [{"runs": [("THE LOUD GOD", {"size": 54, "bold": True, "color": ACCENT})]}])
tx(s, 0.9, 3.7, 7.0, 1.6, [
    {"runs": [("He teaches shouting. Mocking.", {"size": 20, "color": WHITE})]},
    {"runs": [("\u201cBecause I said so.\u201d", {"size": 20, "color": WHITE})], "sb": 4},
    {"runs": [("Every interruption makes him glow brighter.", {"size": 15, "color": ONDARK})], "sb": 12},
])
bunny(s, 10.5, 6.3, 3.7)
note(s, "Cold open part 2. Reveal (or place) the pink glowing plastic bunny on your desk NOW. Deadpan: "
        "\u201cBehold the Loud God.\u201d The cheaper the bunny, the harder they laugh \u2014 hold the stare.")

# ── 4 · Cold open 3 — The Split (room map) ───────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 10.0, 0.7, [{"runs": [("The kingdoms split.", {"size": 36, "bold": True})]}])
tx(s, 0.55, 1.25, 12.2, 0.5, [{"runs": [("Two walls of this room. Every decree, you take a stand.", {"size": 16, "color": MUTED})]}])
bw_, by_ = 4.3, 2.1
b1 = box(s, 0.7, by_, bw_, 2.5, fill=ACCENT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
b2 = box(s, 0.7 + bw_ + 0.05, by_, 3.2, 2.5, fill=MUTED, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
b3 = box(s, 0.7 + bw_ + 3.3, by_, bw_, 2.5, fill=DARK, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
for b, l1, l2 in [(b1, "KINGDOM OF", "YEAHTOTALLY"), (b2, "THE", "MID"), (b3, "REPUBLIC OF", "NOWAY")]:
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = l1
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(15), True, WHITE
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(2)
    r2 = p2.add_run(); r2.text = l2
    r2.font.name, r2.font.size, r2.font.bold, r2.font.color.rgb = FONT, Pt(26), True, WHITE
    p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER; p3.space_before = Pt(8)
    r3 = p3.add_run(); r3.text = {"YEAHTOTALLY": "believes first \u2014 defends with reasons",
                                  "MID": "undecided \u2014 and it costs you",
                                  "NOWAY": "doubts first \u2014 demands evidence"}[l2]
    r3.font.name, r3.font.size, r3.font.bold, r3.font.color.rgb = FONT, Pt(12), False, WHITE
tx(s, 0.7, 5.0, 4.3, 0.8, [{"align": PP_ALIGN.CENTER, "runs": [("\u2190  that wall", {"size": 15, "bold": True, "color": ACCENT})]}])
tx(s, 5.05, 5.0, 3.2, 0.8, [{"align": PP_ALIGN.CENTER, "runs": [("the middle strip", {"size": 15, "bold": True, "color": MUTED})]}])
tx(s, 8.3, 5.0, 4.3, 0.8, [{"align": PP_ALIGN.CENTER, "runs": [("the other wall  \u2192", {"size": 15, "bold": True, "color": DARK})]}])
tx(s, 0.7, 6.2, 12.0, 0.6, [{"align": PP_ALIGN.CENTER,
                             "runs": [("Citizens of the Mid pay a toll: say out loud what would convince you.", {"size": 15, "italic": True, "color": MUTED})]}])
note(s, "Cold open part 3. Physically point: \u201cYeahTotally holds THAT wall. NoWay holds the other. The middle strip "
        "is the Borderlands \u2014 the Mid \u2014 and the Mid pays a toll: say out loud what would convince you.\u201d")

# ── 5 · Cold open 4 — The Mission ────────────────────────────────────────────
s = slide(dark=True)
tx(s, 0.9, 0.95, 9.0, 0.4, [{"runs": [("YOUR MISSION", {"size": 13, "bold": True, "color": ACCENT})]}])
tx(s, 0.86, 1.6, 10.6, 1.7, [{"runs": [("Fill the Ledger of Moved Minds before the hour ends.", {"size": 34, "bold": True, "color": WHITE})]}])
tx(s, 0.9, 3.7, 9.6, 2.4, [
    {"runs": [("You are delegates \u2014 the first 28 ever trusted with both borders.", {"size": 17, "color": ONDARK})]},
    {"runs": [("Defect when a better reason finds you \u2014 and name what moved you.", {"size": 17, "color": ONDARK})], "sb": 10},
    {"runs": [("The Arbiter reads the decrees. The Arbiter never argues.", {"size": 17, "color": ONDARK})], "sb": 10},
])
tx(s, 0.9, 6.5, 9.0, 0.4, [{"runs": [("\u201cLEDGER OF MOVED MINDS\u201d goes on the board \u2014 watch it fill.", {"size": 13, "color": ONDARK})]}])
bunny(s, 12.0, 7.1, 1.5)
note(s, "Cold open part 4. Write LEDGER OF MOVED MINDS on the board now if you haven't. Envelope bit (optional): "
        "open the sealed briefing, glance at it gravely, and read Decree One.")

# ── 6 · Summit Law ───────────────────────────────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 8.0, 0.7, [{"runs": [("Summit law", {"size": 36, "bold": True})]}])
rules_rows(s, [
    ("Attack the argument \u2014 never the arguer.", "No names, no eye-rolls, no \u201cthat\u2019s dumb.\u201d The idea takes the hit."),
    ("Reasons, not volume.", "The Loud God feeds on shouting, mocking, and \u201cbecause I said so.\u201d Reasons starve him."),
    ("Every claim needs a because.", "An opinion without a reason is just noise. Say the because."),
    ("Defect with honour.", "Moving is the noblest act here \u2014 but you must name the argument that moved you."),
    ("Silent feet.", "Move on the countdown (3\u20112\u20111). Talking starts only once partners are set."),
])
box(s, 10.4, 1.5, 2.38, 4.9, fill=TINT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
tx(s, 10.68, 1.85, 1.85, 4.3, [
    {"runs": [("SAY IT LIKE A DELEGATE", {"size": 14, "bold": True, "color": ACCENT})]},
    {"runs": [("\u201cI see it differently because \u2026\u201d", {"size": 13.5, "color": TEXT})], "sb": 12},
    {"runs": [("\u201cThat\u2019s a fair point, but \u2026\u201d", {"size": 13.5, "color": TEXT})], "sb": 10},
    {"runs": [("\u201cWhat\u2019s your evidence?\u201d", {"size": 13.5, "color": TEXT})], "sb": 10},
    {"runs": [("\u201cYou\u2019re right that \u2026 , but \u2026\u201d", {"size": 13.5, "color": TEXT})], "sb": 10},
])
bunny(s, 12.55, 7.25, 0.85, glow=False)
note(s, "Read the five laws fast. Point at the bunny on law 2. Rehearse 'silent feet' once during the first training decree.")

# ── 7 · How a decree works ───────────────────────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 10.0, 0.7, [{"runs": [("How a decree works", {"size": 36, "bold": True})]}])
steps = [
    ("STAND", "The Arbiter reads the decree. 3\u20112\u20111, silent feet \u2014 pick your wall.", "\u224830 s"),
    ("FOLD", "Both ends turn and walk inward. Negotiate with the delegate you meet.", "\u224820 s"),
    ("NEGOTIATE", "60 seconds each way. Reasons, not volume. Starters on the wall.", "60 s"),
    ("VOICES", "Two or three delegates speak to the whole summit. Everyone listens.", "2 min"),
    ("DEFECT", "Persuaded? Cross the floor \u2014 and name the argument that moved you.", "tally"),
]
sw_, sgap_, sy_, sh_ = 2.25, 0.32, 1.7, 2.6
sx_ = (W - (5 * sw_ + 4 * sgap_)) / 2
for i, (name, desc, t) in enumerate(steps):
    x = sx_ + i * (sw_ + sgap_)
    b = box(s, x, sy_, sw_, sh_, fill=PANEL, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    tf = b.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.14)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = name
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(17), True, ACCENT
    p2 = tf.add_paragraph(); p2.space_before = Pt(5)
    r2 = p2.add_run(); r2.text = desc
    r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(11.5), TEXT
    p3 = tf.add_paragraph(); p3.space_before = Pt(7)
    r3 = p3.add_run(); r3.text = t
    r3.font.name, r3.font.size, r3.font.bold, r3.font.color.rgb = FONT, Pt(12), False, MUTED
    if i < 4:
        arrow(s, x + sw_ + 0.04, sy_ + sh_ / 2, x + sw_ + sgap_ - 0.04, sy_ + sh_ / 2)
box(s, 0, 4.75, W, 0.95, fill=TINT)
tx(s, 1.0, 4.75, 11.3, 0.95, [{"align": PP_ALIGN.CENTER,
                               "runs": [("One full decree \u2248 4 minutes. Seven decrees fill the hour \u2014 every defector goes on the Ledger.",
                                         {"size": 16, "italic": True, "color": TEXT})]}], anchor=MSO_ANCHOR.MIDDLE)
tx(s, 1.0, 6.05, 11.3, 0.5, [{"runs": [("Referee signal: when you hear \u201cDECREE!\u201d \u2014 stand, then fold. That\u2019s the whole system.",
                                        {"size": 15, "color": MUTED})]}])
note(s, "Walk the cycle once, slowly. Then run the two training decrees straight away \u2014 the mechanics stick by doing, not by slide.")

# ── 8-13 · Decrees in play order ─────────────────────────────────────────────
decrees = [
    ("I", "TRAINING DECREE I", "Pineapple belongs on pizza.",
     "Rehearsal decree \u2014 learn STAND, FOLD, NEGOTIATE, VOICES, DEFECT before anything matters."),
    ("II", "TRAINING DECREE II", "Cats are better pets than dogs.",
     FOOT),
    ("III", "DECREE OF THE REALM I", "Homework should be banned in every grade.",
     FOOT),
    ("IV", "DECREE OF THE REALM II", "School should start at 10 a.m.",
     FOOT),
    ("V", "DECREE OF THE REALM III", "Video games should be an official school sport.",
     FOOT),
    ("VI", "A WEIGHTIER MATTER", "Zoos do more good than harm.",
     FOOT + "  \u00b7  reasons and examples, not feelings alone"),
]
for num, tag, st, foot in decrees:
    s = decree(num, tag, st, foot)
    note(s, "Full cycle (\u22484 min): STAND \u2192 FOLD \u2192 NEGOTIATE 60s \u2192 VOICES (2\u20133 delegates) \u2192 DEFECT + Ledger tally. "
            "If 90% land on one wall, swap in a spare decree without ceremony.")

# ── 14 · Double Agents ───────────────────────────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 10.0, 0.7, [{"runs": [("Double Agents", {"size": 36, "bold": True})]}])
tx(s, 0.55, 1.25, 12.0, 0.5, [{"runs": [("One decree. Everyone argues the OTHER kingdom\u2019s side.", {"size": 16, "color": MUTED})]}])
rules_rows(s, [
    ("Walk to the other wall.", "You are now a citizen there. Swear temporary loyalty \u2014 this is a spy mission."),
    ("Steal their best argument.", "Use the intelligence you overheard during negotiations \u2014 their strongest reason, now yours."),
    ("Deliver it like you believe it.", "Best double-agent performance earns the crown. Conviction is the costume."),
], y0=2.0, rh=0.95, textw=11.3)
box(s, 0, 5.35, W, 1.0, fill=TINT)
tx(s, 1.0, 5.35, 11.3, 1.0, [{"align": PP_ALIGN.CENTER,
                              "runs": [("Debrief: how did it feel to argue what you don\u2019t believe?",
                                        {"size": 17, "italic": True, "color": TEXT})]}], anchor=MSO_ANCHOR.MIDDLE)
note(s, "Grade-7 devil's advocate: easier and funnier than 'invent the other side' \u2014 they borrow an argument they actually heard. "
        "2\u20133 performances, then the debrief question.")

# ── 15 · Double agent decree ─────────────────────────────────────────────────
s = decree("VII", "DOUBLE AGENTS \u00b7 ARGUE THE OTHER KINGDOM", "Homework should be banned in every grade.",
           "Yes \u2014 again. Your kingdom has switched. Steal their best argument and make it sing.")
note(s, "The reprise decree: they fought this 20 minutes ago, so every delegate already knows both arsenals. Run 60-second "
        "negotiations, then 2\u20133 performances.")

# ── 16 · Delegates' Own Laws ─────────────────────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 10.0, 0.7, [{"runs": [("The Delegates\u2019 Own Laws", {"size": 36, "bold": True})]}])
tx(s, 0.55, 1.3, 8.6, 0.6, [{"runs": [("Write a decree the whole summit must take a stand on.", {"size": 16, "color": MUTED})]}])
tx(s, 0.55, 2.1, 8.6, 0.7, [{"runs": [("school rules  \u00b7  food  \u00b7  weekends  \u00b7  technology  \u00b7  sports  \u00b7  shows & games",
                                       {"size": 17, "bold": True, "color": ACCENT})]}])
for i, (h, g) in enumerate([
    ("Write it on a slip. One sentence.", "Statements, not questions \u2014 something to agree or disagree with."),
    ("The Arbiter vets in silence.", "Three or four become law and get read aloud."),
    ("Authors stay silent.", "Watch the kingdoms take your idea seriously. That\u2019s the reward."),
]):
    y = 3.1 + i * 1.05
    tx(s, 0.55, y - 0.03, 0.75, 0.6, [{"runs": [(str(i + 1), {"size": 26, "bold": True, "color": ACCENT})]}])
    tx(s, 1.45, y, 8.4, 0.95, [
        {"runs": [(h, {"size": 19, "bold": True})]},
        {"runs": [(g, {"size": 13.5, "color": MUTED})], "sb": 3},
    ])
box(s, 10.4, 1.5, 2.38, 4.4, fill=TINT, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
tx(s, 10.68, 1.85, 1.85, 3.9, [
    {"runs": [("TWO LAWS OF LAWMAKING", {"size": 14, "bold": True, "color": ACCENT})]},
    {"runs": [("1.  No naming classmates.", {"size": 14, "color": TEXT})], "sb": 14},
    {"runs": [("2.  Nothing about anyone\u2019s body or identity.", {"size": 14, "color": TEXT})], "sb": 10},
    {"runs": [("Vetoes are silent and final \u2014 no debate.", {"size": 12.5, "color": MUTED})], "sb": 14},
])
note(s, "90 seconds to write. Screen every slip in one pass \u2014 read out the ones that split the room.")

# ── 16b · Hidden: type-in student decree ─────────────────────────────────────
s = decree("D", "A DELEGATE\u2019S OWN LAW \u00b7 READ ALOUD TO THE SUMMIT",
           "[ type the delegate\u2019s decree here \u2014 double-click this text ]",
           "Hidden slide. During the Own Laws round, type a chosen slip here and project it.")
note(s, "Hidden slide. When a student decree deserves the projector, unhide this slide, double-click the statement "
        "text, and type their sentence verbatim (fix nothing \u2014 read their words). Reuse it for Double Agents "
        "or as the Decree of the Hour reveal.")

# ── 17 · Ledger & the Curse ──────────────────────────────────────────────────
s = slide()
tx(s, 0.55, 0.5, 10.0, 0.7, [{"runs": [("The Ledger & the Curse", {"size": 36, "bold": True})]}])
tx(s, 0.55, 1.25, 12.0, 0.5, [{"runs": [("Ten minutes left. Time to find out if the Loud God starves.", {"size": 16, "color": MUTED})]}])
rules_rows(s, [
    ("Count the moved minds.", "Walk the Ledger: which decree defected the most delegates?"),
    ("Crown the Decree of the Hour.", "The statement that moved the class \u2014 was it logic, evidence, or an example?"),
    ("Confess to the Ledger.", "Exit slip: one argument that almost moved you \u2014 and what stopped it."),
], y0=2.1, rh=1.0, textw=11.3)
box(s, 0, 5.55, W, 1.0, fill=TINT)
tx(s, 1.0, 5.55, 11.3, 1.0, [{"align": PP_ALIGN.CENTER,
                              "runs": [("Then the Arbiter declares the curse broken. It always breaks \u2014 the oath decides, not the tally.",
                                        {"size": 16, "italic": True, "color": TEXT})]}], anchor=MSO_ANCHOR.MIDDLE)
note(s, "Collect exit slips on the way out \u2014 the 'almost moved' answers tell you which ILT project hooks this class.")

# ── 18 · The Oath ────────────────────────────────────────────────────────────
s = slide(dark=True)
tx(s, 0.9, 0.85, 9.0, 0.4, [{"runs": [("THE OATH OF THE MOVED MIND", {"size": 13, "bold": True, "color": ACCENT})]}])
lines = [
    ("Arbiter:", "Do you swear to argue with reasons, not volume?", "We swear."),
    ("Arbiter:", "To attack the argument \u2014 never the arguer?", "We swear."),
    ("Arbiter:", "And if a better reason finds you \u2014 will you move without shame?", "I will move."),
]
y = 1.75
for who, q, a in lines:
    tx(s, 0.9, y, 11.5, 0.9, [{"runs": [(who + "  ", {"size": 18, "bold": True, "color": ONDARK}),
                                        (q, {"size": 20, "bold": True, "color": WHITE})]}])
    tx(s, 0.9, y + 0.95, 11.5, 0.55, [{"runs": [(a, {"size": 22, "bold": True, "italic": True, "color": ACCENT})]}])
    y += 1.65
tx(s, 0.9, 6.85, 11.0, 0.4, [{"runs": [("Call and response. Last two minutes of the hour.", {"size": 14, "color": ONDARK})]}])
note(s, "Teach the response pattern in 10 seconds ('I ask, you answer \u2014 loudly, but with reasons'). Deliver the last line, "
        "declare the curse broken, and let them file out past the bunny.")

# ── Hidden spares ────────────────────────────────────────────────────────────
spares = [
    ("A", "SPARE \u00b7 TRAINING DECREE", "Winter is the best season in Nova Scotia."),
    ("B", "SPARE \u00b7 TRAINING DECREE", "The movie is never as good as the book."),
    ("C", "SPARE \u00b7 DECREE OF THE REALM", "Exams should always be open-book."),
    ("D", "SPARE \u00b7 DECREE OF THE REALM", "Students should be allowed to redo any test they fail."),
    ("E", "SPARE \u00b7 A WEIGHTIER MATTER", "It\u2019s never OK to lie \u2014 even to be kind."),
    ("F", "SPARE \u00b7 A WEIGHTIER MATTER", "Professional athletes are paid too much."),
    ("G", "SPARE \u00b7 A WEIGHTIER MATTER", "Junk food should be banned from school cafeterias."),
]
for letter, tag, st in spares:
    s = decree(letter, tag, st, FOOT + "  \u00b7  hidden slide \u2014 unhide to swap in", hidden=True, letter=letter)
    note(s, "Hidden slide. Unhide and print if a decree lands 90% on one wall.")

prs.save(r"F:\Antigravity\simroom\Github Repos\bihi2027\ILT\oneday-06-summit-deck.pptx")
print("saved", len(prs.slides._sldIdLst), "slides")
