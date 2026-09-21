import glob, os, json, re, zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

folder = r'F:\Antigravity\simroom\Github Repos\bihi2027\Citizenship 9\WHERE\903'
out = {}

def walk_shape(shape, out_list):
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            walk_shape(sub, out_list)
        return
    if getattr(shape, 'has_table', False) and shape.has_table:
        rows = []
        for r in shape.table.rows:
            rows.append(' / '.join(c.text.strip().replace('\n', ' ') for c in r.cells))
        out_list.append('[TABLE] ' + ' || '.join(rows))
        return
    if shape.has_text_frame:
        t = shape.text_frame.text.strip()
        if t:
            out_list.append(t.replace('\n', ' | '))
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        out_list.append('[PICTURE]')

# pptx files
for f in sorted(glob.glob(os.path.join(folder, '*.pptx'))):
    base = os.path.basename(f)
    try:
        prs = Presentation(f)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                walk_shape(shape, texts)
            notes = ''
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()[:200]
            slides.append({'slide': i, 'texts': texts, 'notes': notes})
        out[base] = slides
        print(f'\n=== {base}')
        for s in slides:
            print(f"  [s{s['slide']}] {' || '.join(s['texts'])[:350]}")
    except Exception as e:
        print(f'=== {base} :: ERROR {e}')

# docx files
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for f in sorted(glob.glob(os.path.join(folder, '*.docx'))):
    base = os.path.basename(f)
    try:
        with zipfile.ZipFile(f) as z:
            xml = z.read('word/document.xml')
        root = ET.fromstring(xml)
        paras = []
        for p in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
            t = ''.join(n.text or '' for n in p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')).strip()
            if t:
                paras.append(t)
        out[base] = paras
        print(f'\n=== {base}')
        for p in paras[:40]:
            print('  ', p[:250])
    except Exception as e:
        print(f'=== {base} :: ERROR {e}')

json.dump(out, open(os.path.join(folder, '_extracted.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('\nsaved _extracted.json')
